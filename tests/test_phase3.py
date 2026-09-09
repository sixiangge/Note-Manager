"""Offline Phase 3 regressions. Optional integrations skip in the base environment."""

from dataclasses import replace
from importlib.util import find_spec
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from src.app_settings import normalized_settings
from src.models import ExternalDocument, Note
from src.oracle import _analyze_discrepancies, _build_teach_prompt, teach_session, TeachingCancelled
from src.parsers.external import chunk_text, parse_external
from src.rag.retriever import VectorRetriever, load_notes, text_vector
from src.teaching_config import TeachingConfig

HAS_CHROMA = find_spec("chromadb") is not None
HAS_PARSERS = all(find_spec(name) for name in ("pypdf", "pptx", "docx2txt"))


class TemporaryCase(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)

    def note(self, name="极限", subject="微积分", tags="[]", kind="exam"):
        path = self.root / (name + ".md")
        path.write_text(f"---\ntitle: {name}\nsubject: {subject}\nchapter: 第一章\n"
                        f"tags: {tags}\nnote_type: {kind}\nexam_freq: 2\n---\n极限的定义：epsilon 和 delta。",
                        encoding="utf-8")
        return path


class TestTeachingSettings(TemporaryCase):
    def test_defaults_disabled_and_no_credentials(self):
        settings = normalized_settings({"OPENAI_API_KEY": "fake", "teach_api_key": "fake",
                                        "teach_base_url": "https://user:secret@example.org"}, self.root)
        self.assertEqual(settings["teach_provider"], "disabled")
        self.assertNotIn("fake", json.dumps(settings))
        self.assertNotIn("secret", settings["teach_base_url"])
        with self.assertRaises(ValueError):
            TeachingConfig.from_settings(settings).validate()

    def test_config_bounds_and_urls(self):
        valid = TeachingConfig(provider="local", model="demo")
        valid.validate()
        for changes in ({"base_url": "http://example.com"}, {"timeout": 1},
                        {"base_url": "https://example.com?key=fake"}, {"top_k": 0},
                        {"allowed_types": ".exe"}, {"model": ""}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                replace(valid, **changes).validate()
        with patch.dict("os.environ", {}, clear=True), self.assertRaises(ValueError):
            replace(valid, provider="api").validate()

    def test_settings_roundtrip_non_secret(self):
        from src.storage import save_app_settings, load_app_settings
        settings = normalized_settings({"teach_provider": "local", "teach_model": "demo",
                                        "teach_timeout": 999, "teach_top_k": 7}, self.root)
        database = self.root / "state.db"
        save_app_settings(database, settings)
        restored = normalized_settings(load_app_settings(database), self.root)
        self.assertEqual(restored, settings)
        self.assertEqual(restored["teach_timeout"], 300)


class TestExternalText(TemporaryCase):
    def test_chunk_overlap_and_limits(self):
        value = "abcdefghijklmnopqrstuvwxyz"
        chunks = chunk_text(value, 10)
        self.assertEqual(chunks[:2], [value[:10], value[9:19]])
        self.assertEqual(chunk_text(" \n "), [])
        with self.assertRaises(ValueError):
            chunk_text("x", 0)
        with self.assertRaises(ValueError):
            chunk_text("x" * 2_000_001)

    def test_txt_cache_change_and_corruption(self):
        path = self.root / "中文 课件.txt"
        path.write_text("极限的定义" * 30, encoding="utf-8-sig")
        original = path.read_bytes()
        first = parse_external(path, 100)
        self.assertEqual(list(self.root.iterdir()), [path])
        cache = self.root / "cache"
        self.assertEqual(parse_external(path, 100, cache_dir=cache).chunks, first.chunks)
        cached = next(cache.iterdir())
        cached.write_text("invalid json", encoding="utf-8")
        self.assertEqual(parse_external(path, 100, cache_dir=cache).chunks, first.chunks)
        self.assertEqual(path.read_bytes(), original)
        path.write_text("不同的定义", encoding="utf-8")
        self.assertNotEqual(parse_external(path, cache_dir=cache).doc_id, first.doc_id)

    def test_reject_empty_unsupported_large(self):
        path = self.root / "资料.txt"
        path.write_text("", encoding="utf-8")
        with self.assertRaises(ValueError):
            parse_external(path)
        path.write_text("text", encoding="utf-8")
        with self.assertRaises(ValueError):
            parse_external(path, allowed_types=".pdf")
        path.write_bytes(b"x" * (1024 * 1024 + 1))
        with self.assertRaises(ValueError):
            parse_external(path, max_file_mb=1)

    def test_note_filters_and_bad_file_isolation(self):
        wanted = self.note()
        self.note("示例", tags="[示例]")
        self.note("别的科目", subject="数据库")
        self.note("考研", kind="postgraduate")
        (self.root / "坏笔记.md").write_text("bad", encoding="utf-8")
        warnings = []
        notes = load_notes(self.root, "微积分", "exam", warnings=warnings)
        self.assertEqual([n.file_path for n in notes], [wanted])
        self.assertEqual(len(warnings), 1)
        self.assertEqual(len(load_notes(self.root, include_samples=True)), 4)


@unittest.skipUnless(HAS_PARSERS, "Install requirements-phase3.txt for parser integration tests")
class TestOfficePdfParsers(TemporaryCase):
    def test_pdf_text_and_empty_pages(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        path = self.root / "教材.pdf"
        writer = PdfWriter()
        writer.add_blank_page(500, 500)
        page = writer.add_blank_page(500, 500)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
            DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 10 100 Td (limit definition) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(path)
        document = parse_external(path)
        self.assertIn("limit definition", document.chunks[0])
        self.assertIn("第 2 页", document.chunks[0])
        empty = PdfWriter()
        empty.add_blank_page(100, 100)
        empty.write(path)
        with self.assertRaisesRegex(ValueError, "OCR"):
            parse_external(path)

    def test_pdf_password(self):
        from pypdf import PdfWriter
        path = self.root / "加密.pdf"
        writer = PdfWriter()
        writer.add_blank_page(100, 100)
        writer.encrypt("password")
        writer.write(path)
        with self.assertRaisesRegex(ValueError, "加密"):
            parse_external(path)

    def test_pptx_text_table_group_and_cache_type(self):
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "极限课件"
        table = slide.shapes.add_table(1, 1, 0, Inches(1), Inches(3), Inches(1)).table
        table.cell(0, 0).text = "表格中的定义"
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "组合形状"
        path = self.root / "课件.pptx"
        presentation.save(path)
        document = parse_external(path, cache_dir=self.root / "cache")
        self.assertTrue(all(s in document.chunks[0] for s in ("极限课件", "表格中的定义", "组合形状", "幻灯片 1")))
        self.assertEqual(parse_external(path, cache_dir=self.root / "cache").doc_type, "PPT")

    def test_docx_paragraph_and_table(self):
        path = self.root / "教材.docx"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>极限的定义</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r><w:t>表格内容</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>')
        chunks = parse_external(path).chunks
        self.assertIn("极限的定义", chunks[0])
        self.assertIn("表格内容", chunks[0])


@unittest.skipUnless(HAS_CHROMA, "Install requirements-phase3.txt for Chroma integration tests")
class TestRetrievalAndTeaching(TemporaryCase):
    def test_separate_collections_isolation_and_cleanup(self):
        import chromadb
        from chromadb.config import Settings
        client = chromadb.EphemeralClient(Settings(anonymized_telemetry=False))
        before = {c.name for c in client.list_collections()}
        with VectorRetriever(client=client) as first, VectorRetriever(client=client) as second:
            first.index_notes([Note("极限定义", "微积分", "第一章", body="极限定义 epsilon delta", file_path=self.root / "note.md")])
            first.index_external([ExternalDocument("external", self.root / "课件.txt", chunks=["极限定义 epsilon delta"])])
            self.assertEqual(first.query("极限定义", source_type="external")[0]["source_type"], "external")
            self.assertEqual(first.query("极限定义", source_type="note")[0]["source_type"], "note")
            self.assertEqual(second.query("极限定义"), [])
            self.assertEqual(first.query("-***"), [])
            first.clear_external()
            self.assertEqual(first.query("极限定义", source_type="external"), [])
            self.assertTrue(first.query("极限定义", source_type="note"))
        self.assertEqual({c.name for c in client.list_collections()}, before)

    def test_teaching_end_to_end_with_fake_model(self):
        path = self.note()
        original = path.read_bytes()
        attachment = self.root / "课件.txt"
        attachment.write_text("极限的定义 epsilon delta 外部正确解释", encoding="utf-8")
        response = {"answer": "根据 [E1] 和 [N1]，解释极限。", "discrepancies": [
            {"kind": "冲突", "message": "定义需要核对", "note_ids": ["N1"], "external_ids": ["E1"]},
            {"kind": "可能遗漏", "message": "补充考点", "external_ids": ["E1"]},
            {"kind": "补充", "message": "笔记补充观点", "note_ids": ["N1"]}]}
        with patch("src.oracle.generate", return_value=response) as model:
            result = teach_session("极限的定义", [str(attachment), str(self.root / "bad.pdf")],
                config=TeachingConfig(provider="local", model="fake"), notes_root=self.root)
        self.assertEqual(result.final_answer, response["answer"])
        self.assertEqual(len(result.discrepancies), 3)
        self.assertTrue(result.warnings)
        self.assertEqual(path.read_bytes(), original)
        request = model.call_args.args[1][-1]["content"]
        self.assertNotIn(str(self.root), request)
        self.assertIn("极限", request)
        self.assertFalse((self.root / "cache").exists())

    def test_cancel_failure_and_empty_do_not_leak_collections(self):
        import chromadb
        from chromadb.config import Settings
        client = chromadb.EphemeralClient(Settings(anonymized_telemetry=False))
        before = {c.name for c in client.list_collections()}
        config = TeachingConfig(provider="local", model="fake")
        with patch("src.oracle.generate") as model:
            result = teach_session("极限", config=config, notes_root=self.root)
            self.assertIn("足够相关", result.final_answer)
            model.assert_not_called()
            with self.assertRaises(TeachingCancelled):
                teach_session("极限", config=config, notes_root=self.root, cancelled=lambda: True)
        self.note()
        with patch("src.oracle.generate", side_effect=RuntimeError("failure")), self.assertRaises(RuntimeError):
            teach_session("极限", config=config, notes_root=self.root)
        self.assertEqual({c.name for c in client.list_collections()}, before)

    def test_all_failed_uploads_do_not_send(self):
        with patch("src.oracle.generate") as model, self.assertRaisesRegex(ValueError, "所有附件"):
            teach_session("极限", [str(self.root / "missing.txt")],
                          config=TeachingConfig(provider="local", model="fake"), notes_root=self.root)
        model.assert_not_called()

    def test_fabricated_answer_reference_rejected(self):
        self.note()
        with patch("src.oracle.generate", return_value={"answer": "依据 [E99]", "discrepancies": []}):
            with self.assertRaisesRegex(ValueError, "未提供的来源"):
                teach_session("极限定义", config=TeachingConfig(provider="local", model="fake"), notes_root=self.root)


class TestEvidenceValidation(unittest.TestCase):
    def test_invalid_citations_rejected(self):
        for finding in ({"kind": "冲突", "message": "冲突", "note_ids": ["N9"], "external_ids": ["E1"]},
                        {"kind": "冲突", "message": "冲突", "note_ids": [], "external_ids": ["E1"]}):
            with self.assertRaises(ValueError):
                _analyze_discrepancies([{"id": "N1"}], [{"id": "E1"}], [finding])

    def test_vectors_are_deterministic_and_ignore_punctuation(self):
        self.assertEqual(text_vector("E-R模型"), text_vector("E-R模型"))
        self.assertFalse(any(text_vector("- ! $")))
        self.assertAlmostEqual(sum(v*v for v in text_vector("极限的定义")), 1)
