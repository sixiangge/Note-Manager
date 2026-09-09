"""Extract slide text, tables and grouped shapes without running Office."""

from pathlib import Path

from src.models import ExternalDocument
from src.parsers.external import document_from_text


def _shape_text(shape: object) -> str:
    if hasattr(shape, "shapes"):
        return "\n".join(_shape_text(child) for child in shape.shapes)
    if getattr(shape, "has_table", False):
        return "\n".join(" | ".join(cell.text for cell in row.cells)
                         for row in shape.table.rows)
    return getattr(shape, "text", "")


def parse_pptx(file_path: Path, chunk_size: int = 500) -> ExternalDocument:
    """Extract labelled slides.

    Args:
        file_path: PowerPoint path.
        chunk_size: Maximum characters per chunk.

    Returns:
        Parsed external document.
    """
    from pptx import Presentation

    path = Path(file_path)
    slides = []
    for number, slide in enumerate(Presentation(path).slides, 1):
        text = "\n".join(_shape_text(shape) for shape in slide.shapes).strip()
        if text:
            slides.append(f"[幻灯片 {number}]\n{text}")
    return document_from_text(path, "\n\n".join(slides), chunk_size)
